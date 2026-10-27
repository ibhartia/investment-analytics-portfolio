"""Tie-out tests: the Python pipelines must reproduce the Excel workbooks exactly.

The workbooks ship with cached values (recalculated in LibreOffice), so these
tests read them with data_only=True and compare against the Python output.
Run from the repo root:  pytest -q
"""
import importlib.util
import math
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


sc = _load("scorecard", ROOT / "earnings-call-sentiment/src/scorecard.py")
ev = _load("exit_valuation", ROOT / "vc-exit-valuation/src/exit_valuation.py")


# ---------------- Earnings sentiment ----------------
@pytest.fixture(scope="module")
def scorecard():
    w = sc.load_weights()
    s = sc.load_scores()
    scored = sc.add_composite(s, w)
    return w, s, scored, sc.company_summary(scored)


def test_scores_clean(scorecard):
    w, s, _, _ = scorecard
    assert sc.validate(s, w) == []


def test_validation_catches_bad_rows(scorecard):
    w, s, _, _ = scorecard
    bad = s.copy()
    bad.loc[0, "mgmt_tone"] = 7
    bad.loc[1, "call_date"] = pd.NaT
    issues = sc.validate(bad, w * 1.1)
    assert any("Weights sum" in i for i in issues)
    assert any("out of range" in i for i in issues)
    assert any("invalid or missing call date" in i for i in issues)


def test_composites_match_workbook(scorecard):
    _, _, scored, _ = scorecard
    ws = openpyxl.load_workbook(ROOT / "earnings-call-sentiment/workbook/Earnings_Sentiment_Scorecard.xlsx",
                                data_only=True)["All Scores"]
    xl = {(ws.cell(r, 1).value, ws.cell(r, 3).value): ws.cell(r, 11).value for r in range(3, 52)}
    for r in scored.itertuples():
        assert math.isclose(r.composite, xl[(r.ticker, r.quarter)], abs_tol=1e-9), (r.ticker, r.quarter)


def test_company_summary_matches_workbook(scorecard):
    *_, summary = scorecard
    ws = openpyxl.load_workbook(ROOT / "earnings-call-sentiment/workbook/Earnings_Sentiment_Scorecard.xlsx",
                                data_only=True)["Company Summary"]
    s = summary.set_index("ticker")
    for r in range(3, 13):
        t = ws.cell(r, 1).value
        assert ws.cell(r, 3).value == s.loc[t, "quarters"]
        assert math.isclose(ws.cell(r, 10).value, s.loc[t, "avg_composite"], abs_tol=0.005)
        assert math.isclose(ws.cell(r, 11).value, s.loc[t, "trend"], abs_tol=0.005)


def test_custom_weights_change_composite(scorecard):
    w, s, scored, _ = scorecard
    w2 = pd.Series(0.0, index=w.index)
    w2["fwd_guidance"] = 1.0
    alt = sc.add_composite(s, w2)
    assert (alt["composite"] == s["fwd_guidance"]).all()


# ---------------- Exit valuation ----------------
@pytest.fixture(scope="module")
def valuation():
    a = ev.Assumptions.from_csv()
    p, c, l = ev.load_inputs()
    return ev.value_portfolio(p, c, l, a), a


@pytest.fixture(scope="module")
def workbook():
    return openpyxl.load_workbook(ROOT / "vc-exit-valuation/workbook/Exit_Valuation_Framework_SAMPLE.xlsx",
                                  data_only=True)


COLS = {  # Excel column -> Python column (Valuation Model)
    "F": "n_exit", "G": "exit_val_exit_comps", "H": "n_private", "K": "exit_val_private_comps",
    "L": "exit_value", "P": "p_used", "S": "exit_year", "T": "gross_exit_to_fund", "U": "risk_adj",
    "V": "risk_adj_exit_basis", "W": "risk_adj_private_basis", "X": "pv_risk_adj", "Z": "exp_moic",
    "AA": "legacy_exit",
}


def _close(a, b):
    if a in (None, "") and (b is None or (isinstance(b, float) and math.isnan(b))):
        return True
    return math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=1e-6)


def test_every_company_matches_workbook(valuation, workbook):
    v, _ = valuation
    ws = workbook["Valuation Model"]
    for i, row in enumerate(v.itertuples()):
        r = 3 + i
        assert ws[f"A{r}"].value == row.company
        for col, field in COLS.items():
            assert _close(ws[f"{col}{r}"].value, getattr(row, field)), (row.company, col)


def test_reserves_match_model_2(valuation, workbook):
    v, _ = valuation
    ws = workbook["Valuation Model 2"]
    for i, row in enumerate(v.itertuples()):
        r = 3 + i
        for col, field in {"AD": "follow_on", "AE": "capital_deployed", "AF": "net_exp_moic"}.items():
            assert _close(ws[f"{col}{r}"].value, getattr(row, field)), (row.company, col)


def test_summary_matches_workbook(valuation, workbook):
    v, _ = valuation
    s = ev.summarize(v)
    ws = workbook["Summary"]
    expected = {4: "follow_on_reserves", 5: "gross_exit_value", 6: "risk_adjusted_value", 7: "total_invested",
                8: "expected_moic", 9: "expected_moic_ex_realized", 10: "pv_risk_adjusted",
                11: "net_expected_moic", 12: "legacy_exit_value"}
    for r, key in expected.items():
        assert _close(ws[f"B{r}"].value, s[key]), key


def test_workbook_totals_internally_consistent(workbook):
    ws = workbook["Valuation Model"]
    for col in "TUVWXY":
        assert _close(ws[f"{col}45"].value, sum(ws[f"{col}{r}"].value or 0 for r in range(3, 45))), col


def test_stage_never_changes_exit_value(valuation):
    """Stage drives probability and timing only: re-staging a company must not move its exit value."""
    v, a = valuation
    p, c, l = ev.load_inputs()
    p2 = p.copy()
    p2.loc[p2["stage"] == "S4", "stage"] = "S7"
    v2 = ev.value_portfolio(p2, c, l, a)
    assert (v["exit_value"].to_numpy() == v2["exit_value"].to_numpy()).all()


def test_no_confidential_identifiers():
    """Guard against real portfolio identifiers leaking into the public repo."""
    banned = ["CSTV", "XXXX", "Prometheus", "Evida", "Femom", "NutriLeap", "Volnay", "Typhoon", "Luminare",
              "Abcuro", "KIMIA", "Satori", "CardioVision", "1Cell", "Tandon", "Wyss"]
    hits = []
    for f in ROOT.rglob("*"):
        if f.is_file() and f.suffix in {".csv", ".md", ".py", ".json"} and f.name != "test_portfolio.py":
            text = f.read_text(errors="ignore")
            hits += [(f.name, b) for b in banned if b in text]
    wb = openpyxl.load_workbook(ROOT / "vc-exit-valuation/workbook/Exit_Valuation_Framework_SAMPLE.xlsx")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    hits += [(f"{ws.title}!{cell.coordinate}", b) for b in banned if b in cell.value]
    assert hits == []
